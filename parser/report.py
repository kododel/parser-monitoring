"""Выгрузка в Excel.

Пять листов: товары, изменения, история цен, динамика с графиком и сводка.
График нужен не для красоты: на встрече с заказчиком динамика цен
убеждает быстрее любой таблицы.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from parser.storage import Storage, to_local

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAX_PRODUCTS_ON_CHART = 5


def _write_header(sheet: Worksheet, titles: list[str], widths: list[int]) -> None:
    for index, title in enumerate(titles, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"


def _build_products_sheet(sheet: Worksheet, products: list) -> None:
    titles = ["№", "Название", "Цена", "Валюта", "В наличии", "Рейтинг", "Первый раз", "Последний раз", "Ссылка"]
    widths = [6, 46, 12, 9, 11, 10, 18, 18, 60]
    _write_header(sheet, titles, widths)

    for row_index, row in enumerate(products, start=2):
        price = row["last_price"]
        currency = row["last_currency"] or "UNKNOWN"
        values = [
            row["id"],
            row["title"],
            price,
            currency,
            "да" if row["last_in_stock"] else "нет",
            row["rating"] or "",
            to_local(row["first_seen_at"]).strftime("%d.%m.%Y %H:%M"),
            to_local(row["last_seen_at"]).strftime("%d.%m.%Y %H:%M"),
            row["url"],
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if col_index in (1, 3, 4, 5, 6) else "left",
            )
        if price is not None:
            sheet.cell(row=row_index, column=3).number_format = "# ##0.00"

    if products:
        sheet.auto_filter.ref = f"A1:I{len(products) + 1}"


def _build_changes_sheet(sheet: Worksheet, changes: list) -> None:
    titles = ["Когда", "Сайт", "Товар", "Было", "Стало", "Разница", "%", "Ссылка"]
    widths = [18, 12, 44, 12, 12, 12, 9, 55]
    _write_header(sheet, titles, widths)

    for row_index, row in enumerate(changes, start=2):
        old_price = row["old_price"]
        new_price = row["new_price"]
        diff = round(new_price - old_price, 2)
        percent = round((new_price - old_price) / old_price * 100, 2) if old_price else 0.0

        values = [
            to_local(row["detected_at"]).strftime("%d.%m.%Y %H:%M"),
            row["site"],
            row["title"],
            old_price,
            new_price,
            diff,
            percent,
            row["url"],
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if col_index in (1, 4, 5, 6, 7) else "left",
            )
        for col_index in (4, 5, 6):
            sheet.cell(row=row_index, column=col_index).number_format = "# ##0.00"
        sheet.cell(row=row_index, column=7).number_format = "0.0"

        # Подешевело — зелёным, подорожало — красным
        colour = "1E7B34" if diff < 0 else "C00000"
        for col_index in (6, 7):
            sheet.cell(row=row_index, column=col_index).font = Font(color=colour, bold=True)

    if changes:
        sheet.auto_filter.ref = f"A1:H{len(changes) + 1}"


def _build_history_sheet(sheet: Worksheet, history: list) -> None:
    titles = ["Товар", "Когда", "Цена", "Валюта", "В наличии"]
    widths = [46, 18, 12, 9, 11]
    _write_header(sheet, titles, widths)

    for row_index, row in enumerate(history, start=2):
        values = [
            row["title"],
            to_local(row["captured_at"]).strftime("%d.%m.%Y %H:%M"),
            row["price"],
            row["currency"],
            "да" if row["in_stock"] else "нет",
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top", horizontal="center" if col_index in (2, 3, 4, 5) else "left"
            )
        sheet.cell(row=row_index, column=3).number_format = "# ##0.00"

    if history:
        sheet.auto_filter.ref = f"A1:E{len(history) + 1}"


def _build_dynamics_sheet(sheet: Worksheet, history: list) -> None:
    """Сводная таблица «замер × товар» и график по ней."""
    by_product: dict[int, dict] = {}
    for row in history:
        entry = by_product.setdefault(
            row["product_id"], {"title": row["title"], "points": {}}
        )
        entry["points"][row["captured_at"]] = row["price"]

    top = sorted(
        by_product.items(), key=lambda item: len(item[1]["points"]), reverse=True
    )[:MAX_PRODUCTS_ON_CHART]

    if not top:
        sheet.cell(row=1, column=1, value="Данных пока нет")
        return

    timestamps = sorted({ts for _, entry in top for ts in entry["points"]})
    if len(timestamps) < 2:
        sheet.cell(row=1, column=1, value="Нужно минимум два замера, чтобы построить график")
        sheet.cell(row=2, column=1, value=f"Сейчас замеров: {len(timestamps)}")
        return

    sheet.cell(row=1, column=1, value="Замер").font = Font(bold=True)
    for col_index, (_, entry) in enumerate(top, start=2):
        cell = sheet.cell(row=1, column=col_index, value=entry["title"][:40])
        cell.font = Font(bold=True)

    for row_index, timestamp in enumerate(timestamps, start=2):
        sheet.cell(
            row=row_index, column=1, value=to_local(timestamp).strftime("%d.%m %H:%M")
        )
        for col_index, (_, entry) in enumerate(top, start=2):
            value = entry["points"].get(timestamp)
            if value is not None:
                sheet.cell(row=row_index, column=col_index, value=value)

    sheet.column_dimensions["A"].width = 16
    for col_index in range(2, len(top) + 2):
        sheet.column_dimensions[get_column_letter(col_index)].width = 18

    chart = LineChart()
    chart.title = "Динамика цен"
    chart.y_axis.title = "Цена"
    chart.x_axis.title = "Замер"
    chart.height = 9
    chart.width = 22

    data = Reference(
        sheet,
        min_col=2,
        max_col=len(top) + 1,
        min_row=1,
        max_row=len(timestamps) + 1,
    )
    categories = Reference(sheet, min_col=1, min_row=2, max_row=len(timestamps) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    sheet.add_chart(chart, f"A{len(timestamps) + 4}")


def _build_summary_sheet(sheet: Worksheet, stats: dict, products: list) -> None:
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 20

    title = sheet.cell(row=1, column=1, value="Сводка по мониторингу")
    title.font = TITLE_FONT
    sheet.merge_cells("A1:B1")

    in_stock = sum(1 for row in products if row["last_in_stock"])
    last_run = stats.get("last_run")

    rows = [
        ("Товаров отслеживается", stats.get("products", 0)),
        ("Всего замеров цены", stats.get("measurements", 0)),
        ("Изменений цены за всё время", stats.get("changes", 0)),
        ("Изменений за 7 дней", stats.get("changes_week", 0)),
        ("Сейчас в наличии", in_stock),
        (
            "Последний сбор",
            to_local(last_run).strftime("%d.%m.%Y %H:%M") if last_run else "ещё не было",
        ),
    ]

    for offset, (label, value) in enumerate(rows, start=3):
        label_cell = sheet.cell(row=offset, column=1, value=label)
        label_cell.border = BORDER
        value_cell = sheet.cell(row=offset, column=2, value=value)
        value_cell.border = BORDER
        value_cell.alignment = Alignment(horizontal="left")


def build_workbook(storage: Storage) -> Workbook:
    """Собирает книгу целиком. Синхронная — вызывается из потока."""
    products = storage.all_products()
    changes = storage.all_changes()
    history = storage.history()

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Товары"
    _build_products_sheet(sheet, products)

    _build_changes_sheet(workbook.create_sheet("Изменения"), changes)
    _build_history_sheet(workbook.create_sheet("История цен"), history)
    _build_dynamics_sheet(workbook.create_sheet("Динамика"), history)
    _build_summary_sheet(workbook.create_sheet("Сводка"), storage.stats(), products)

    return workbook


def export_to_excel(storage: Storage, export_dir: Path) -> Path:
    """Сохраняет выгрузку и возвращает путь к файлу."""
    export_dir = Path(export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)

    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    path = export_dir / f"ceny_{stamp}.xlsx"

    workbook = build_workbook(storage)
    workbook.save(path)

    logger.info("Выгрузка сохранена: %s", path)
    return path
